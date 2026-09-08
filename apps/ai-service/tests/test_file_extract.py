import io
from unittest.mock import patch

import docx

from app.services.file_extract import (
    MIN_PHOTO_BYTES,
    _guess_mime,
    _extract_docx_text,
    extract_photo,
)

# A syntactically valid 1x1 PNG (needed for python-docx's built-in image-header
# parser, which reads real PNG chunks — it doesn't require Pillow).
_TINY_PNG = bytes.fromhex(
    "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c489"
    "0000000a4944415478da63000100000500010d0a2db40000000049454e44ae426082"
)


def _docx_with_picture(image_bytes: bytes) -> bytes:
    document = docx.Document()
    document.add_paragraph("Ada Lovelace")
    document.add_picture(io.BytesIO(image_bytes))
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def test_guess_mime_prefers_file_extension():
    assert _guess_mime("photo.png", b"") == "image/png"
    assert _guess_mime("photo.JPG", b"") == "image/jpeg"
    assert _guess_mime("photo.gif", b"") == "image/gif"


def test_guess_mime_falls_back_to_magic_bytes_when_extension_is_unknown():
    assert _guess_mime("~0~", _TINY_PNG) == "image/png"
    assert _guess_mime(None, b"\xff\xd8\xff\xe0rest") == "image/jpeg"


def test_extract_photo_finds_the_embedded_image_in_a_docx():
    """Regression test: a real .docx with an embedded picture should have
    that picture extracted (as raw bytes + mime type) — the synthetic image
    here is tiny, so MIN_PHOTO_BYTES is patched down to let it through."""
    file_bytes = _docx_with_picture(_TINY_PNG)

    with patch("app.services.file_extract.MIN_PHOTO_BYTES", 10):
        result = extract_photo(
            "resume.docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            file_bytes,
        )

    assert result is not None
    mime_type, image_bytes = result
    assert mime_type == "image/png"
    assert image_bytes == _TINY_PNG


def test_extract_photo_ignores_images_below_the_minimum_size():
    """Small embedded images (decorative icons/logos/lines) shouldn't be
    mistaken for the resume's headshot."""
    file_bytes = _docx_with_picture(_TINY_PNG)

    assert len(_TINY_PNG) < MIN_PHOTO_BYTES
    result = extract_photo(
        "resume.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        file_bytes,
    )

    assert result is None


def test_extract_photo_picks_the_largest_candidate():
    """When multiple images are embedded, the largest one (by byte size) is
    treated as the profile photo — the simplest dependency-free proxy for
    "the headshot" vs. small decorative graphics."""
    small = ("image/png", b"x" * 50)
    large = ("image/jpeg", b"y" * 5000)

    with patch("app.services.file_extract._docx_images", return_value=[small, large]):
        result = extract_photo(
            "resume.docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            b"unused",
        )

    assert result == large


def test_extract_photo_returns_none_for_files_with_no_images():
    file_bytes = docx.Document()
    file_bytes.add_paragraph("No pictures here.")
    buf = io.BytesIO()
    file_bytes.save(buf)

    result = extract_photo(
        "resume.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        buf.getvalue(),
    )

    assert result is None


def test_extract_photo_returns_none_for_unsupported_file_types():
    assert extract_photo("resume.txt", "text/plain", b"plain text") is None


def test_extract_docx_text_includes_table_and_nested_table_content():
    """Regression test: many resume templates lay content out in tables
    (e.g. a two-column layout with a nested table for the sidebar), which
    `document.paragraphs` alone doesn't see."""
    document = docx.Document()
    document.add_paragraph("Ada Lovelace")

    outer_table = document.add_table(rows=1, cols=1)
    inner_table = outer_table.rows[0].cells[0].add_table(rows=1, cols=1)
    inner_table.rows[0].cells[0].paragraphs[0].add_run("Staff Software Engineer")

    buf = io.BytesIO()
    document.save(buf)

    text = _extract_docx_text(buf.getvalue())

    assert "Ada Lovelace" in text
    assert "Staff Software Engineer" in text


def test_extract_docx_text_marks_word_bullet_list_paragraphs_with_a_bullet_prefix():
    """Regression test: Word stores a bulleted paragraph's "•" as numbering
    formatting rather than literal run text, so `paragraph.text` alone loses
    all trace it was a bullet. Extraction should reinsert a "• " marker for
    such paragraphs so downstream mapping (richtext_from_plain) can render
    them as a bulleted list rather than plain paragraphs."""
    document = docx.Document()
    document.add_paragraph("Professional Summary")
    document.add_paragraph(
        "Led AI initiatives across three teams.", style="List Bullet"
    )
    document.add_paragraph("Reduced infrastructure costs by 30%.", style="List Bullet")

    buf = io.BytesIO()
    document.save(buf)

    text = _extract_docx_text(buf.getvalue())
    lines = text.splitlines()

    assert "Professional Summary" in lines
    assert "• Led AI initiatives across three teams." in lines
    assert "• Reduced infrastructure costs by 30%." in lines
