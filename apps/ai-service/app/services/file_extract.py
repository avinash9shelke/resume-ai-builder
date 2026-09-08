"""Extracts raw text and (optionally) an embedded photo from uploaded resume
files (PDF/DOCX).

Feature 1 (Resume Upload and Parsing): the extracted text is fed into the
LangGraph parsing workflow (see app/agents/parser_graph.py). The extracted
photo (if any) is applied directly to the resulting Resume's `picture.url`
by app/api/routes_parse.py — it's deterministic byte extraction, not
something an LLM needs to be involved in.
"""

import io

import docx
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from pypdf import PdfReader

SUPPORTED_CONTENT_TYPES = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
}

# Skip embedded images smaller than this (bytes) when looking for a photo —
# resumes with a headshot also often have tiny decorative icons/logos/lines
# embedded as images, which we don't want to mistake for the profile photo.
MIN_PHOTO_BYTES = 2000


class UnsupportedFileTypeError(ValueError):
    pass


def _resolve_kind(filename: str, content_type: str) -> str:
    kind = SUPPORTED_CONTENT_TYPES.get(content_type)
    if kind is not None:
        return kind
    if filename.lower().endswith(".pdf"):
        return "pdf"
    if filename.lower().endswith(".docx"):
        return "docx"
    raise UnsupportedFileTypeError(f"Unsupported file type: {content_type or filename}")


def extract_text(filename: str, content_type: str, raw_bytes: bytes) -> str:
    kind = _resolve_kind(filename, content_type)
    if kind == "pdf":
        return _extract_pdf_text(raw_bytes)
    return _extract_docx_text(raw_bytes)


def extract_photo(
    filename: str, content_type: str, raw_bytes: bytes
) -> tuple[str, bytes] | None:
    """Returns the largest embedded image (by byte size) as `(mime_type,
    image_bytes)`, or None if the file has no images at or above
    `MIN_PHOTO_BYTES`. Picking the largest is a simple, dependency-free proxy
    for "the headshot" (as opposed to small decorative icons/logos), since a
    resume typically embeds at most one photo."""
    try:
        kind = _resolve_kind(filename, content_type)
    except UnsupportedFileTypeError:
        return None
    try:
        candidates = (
            _pdf_images(raw_bytes) if kind == "pdf" else _docx_images(raw_bytes)
        )
    except (
        Exception
    ):  # noqa: BLE001 - a malformed/unusual file shouldn't break the upload
        return None
    photos = [c for c in candidates if len(c[1]) >= MIN_PHOTO_BYTES]
    if not photos:
        return None
    return max(photos, key=lambda c: len(c[1]))


def _extract_pdf_text(raw_bytes: bytes) -> str:
    reader = PdfReader(io.BytesIO(raw_bytes))
    return "\n".join(page.extract_text() or "" for page in reader.pages).strip()


def _extract_docx_text(raw_bytes: bytes) -> str:
    """Many resume templates use tables (often nested, e.g. for a sidebar/
    two-column layout) to position content rather than plain body paragraphs.
    `document.paragraphs` only sees top-level body paragraphs and silently
    drops anything inside a table cell, so instead walk every `<w:p>` element
    in the document body — this covers top-level paragraphs plus any
    (arbitrarily nested) table cell paragraphs, in reading order."""
    document = docx.Document(io.BytesIO(raw_bytes))
    lines = []
    for p in document.element.body.iter(qn("w:p")):
        paragraph = Paragraph(p, document)
        text = paragraph.text.strip()
        if not text:
            continue
        if _is_list_item(p, paragraph):
            text = f"• {text}"
        lines.append(text)
    return "\n".join(lines).strip()


def _is_list_item(paragraph_element, paragraph: Paragraph) -> bool:
    """True for Word bullet/numbered list paragraphs. Word stores the bullet
    glyph itself as numbering formatting (`<w:numPr>`), not as run text, so
    `paragraph.text` alone loses all trace that a line was a bullet — this
    reinserts a "• " marker so `richtext_from_plain` can map it to a
    bullet-type RichText entry (matching how the resume actually looked)
    instead of a plain paragraph."""
    if paragraph_element.find(".//" + qn("w:numPr")) is not None:
        return True
    style_name = (paragraph.style.name if paragraph.style else "") or ""
    style_name = style_name.lower()
    return "list" in style_name or style_name.endswith("_li")


def _pdf_images(raw_bytes: bytes) -> list[tuple[str, bytes]]:
    reader = PdfReader(io.BytesIO(raw_bytes))
    images: list[tuple[str, bytes]] = []
    for page in reader.pages:
        for image in page.images:
            if image.data:
                images.append((_guess_mime(image.name, image.data), image.data))
    return images


def _docx_images(raw_bytes: bytes) -> list[tuple[str, bytes]]:
    document = docx.Document(io.BytesIO(raw_bytes))
    images: list[tuple[str, bytes]] = []
    for rel in document.part.rels.values():
        if "image" not in rel.reltype:
            continue
        part = rel.target_part
        data = getattr(part, "blob", None)
        if data:
            mime = getattr(part, "content_type", "") or _guess_mime(part.partname, data)
            images.append((mime, data))
    return images


def _guess_mime(name: object, data: bytes) -> str:
    lowered = str(name or "").lower()
    if lowered.endswith(".png"):
        return "image/png"
    if lowered.endswith((".jpg", ".jpeg")):
        return "image/jpeg"
    if lowered.endswith(".gif"):
        return "image/gif"
    if lowered.endswith(".webp"):
        return "image/webp"
    # Fall back to sniffing the magic bytes when the name has no/an unknown extension.
    if data[:8] == b"\x89PNG\r\n\x1a\n":
        return "image/png"
    if data[:3] == b"\xff\xd8\xff":
        return "image/jpeg"
    if data[:6] in (b"GIF87a", b"GIF89a"):
        return "image/gif"
    return "image/jpeg"
